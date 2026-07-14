#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
本地模板库服务（支持网页上传新增图组）。

用法：
    python3 工具/启动模板库.py
然后浏览器打开提示的地址（默认 http://127.0.0.1:8765 ），
在右上角「+ 新增图组」上传图片组。提交后自动存图、写 Excel、重建清单。

依赖：openpyxl   （pip install openpyxl）
仅用标准库 + openpyxl，无需联网。
"""
import os, sys, json, re, base64, subprocess, webbrowser, threading, time
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse

HERE = os.path.dirname(os.path.abspath(__file__))
LIB = os.path.dirname(HERE)
TPL = os.path.join(LIB, "templates")
DATA = os.path.join(LIB, "data")
XLSX = os.path.join(LIB, "元数据录入模板.xlsx")
DICT = os.path.join(DATA, "标签字典-分类目.json")
PORT = int(os.environ.get("PORT", "8765"))
IMG_EXT = (".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp")

try:
    from openpyxl import load_workbook
except ImportError:
    print("！缺少 openpyxl，请先运行:  pip install openpyxl")
    sys.exit(1)


def load_dict():
    return json.load(open(DICT, encoding="utf-8"))


def schema_for(D, cat):
    sc = D["schemas"].get(cat)
    if sc and sc.get("group_level"):
        return cat, sc
    return "通用", D["schemas"]["通用"]


def safe(s):
    return re.sub(r'[\\/:*?"<>|]', "·", str(s)).strip()


def first(v):
    if isinstance(v, list):
        v = v[0] if v else ""
    return re.split(r"[、,，/\s]+", str(v).strip())[0] if v else ""


def age_short(v):
    return re.sub(r"\(.*?\)", "", first(v))


def gen_name(cat, gtags):
    """图组名 = 模特人种·年龄·场景(拍摄形式)·拍摄风格；缺字段则用该类目前几个维度兜底。"""
    parts = [first(gtags.get("模特人种")), age_short(gtags.get("年龄段")),
             first(gtags.get("拍摄形式")), first(gtags.get("拍摄风格"))]
    parts = [p for p in parts if p]
    if not parts:  # 通用类目兜底
        for k in ("拍摄风格", "拍摄场地", "模特风格", "适用场景"):
            if gtags.get(k):
                parts.append(first(gtags[k]))
    return safe("·".join(parts) or "新图组")


def existing_folders(cat):
    d = os.path.join(TPL, cat)
    return set(os.listdir(d)) if os.path.isdir(d) else set()


def unique_name(cat, base):
    used = existing_folders(cat)
    if base not in used:
        return base
    i = 2
    while f"{base}·{i}" in used:
        i += 1
    return f"{base}·{i}"


def existing_ids():
    ids = set()
    if os.path.exists(XLSX):
        wb = load_workbook(XLSX, read_only=True, data_only=True)
        for sn in wb.sheetnames:
            if sn.endswith("图组"):
                ws = wb[sn]; head = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
                if "图组ID" in head:
                    idx = head.index("图组ID")
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if row[idx]:
                            ids.add(str(row[idx]))
        wb.close()
    return ids


def cell(v):
    """把多选列表转成顿号字符串，供 Excel 单元格写入。"""
    if isinstance(v, (list, tuple)):
        return "、".join(str(x) for x in v)
    return "" if v is None else v


def append_excel(cat, gid, name, model, mtype, scene, cover, gtags, images):
    sheet_g = "童装图组" if cat == "童装" else "通用图组"
    sheet_i = "童装图片" if cat == "童装" else "通用图片"
    wb = load_workbook(XLSX)
    if sheet_g not in wb.sheetnames or sheet_i not in wb.sheetnames:
        sheet_g, sheet_i = "通用图组", "通用图片"
    wsg = wb[sheet_g]; headg = [c.value for c in wsg[1]]
    grow = {"类目": cat, "图组ID": gid, "图组名称": name, "模特": model,
            "模特类型": mtype, "场景": scene, "封面文件": cover}
    grow.update(gtags)
    wsg.append([cell(grow.get(h, "")) for h in headg])
    wsi = wb[sheet_i]; headi = [c.value for c in wsi[1]]
    for im in images:
        r = {"图组ID": gid, "图片文件": im["rel"]}
        r.update(im["tags"])
        wsi.append([cell(r.get(h, "")) for h in headi])
    wb.save(XLSX)


def add_group(payload):
    D = load_dict()
    cat = payload.get("category") or "童装"
    gtags = payload.get("group_tags", {})
    images = payload.get("images", [])
    if not images:
        return {"ok": False, "error": "没有图片"}
    name = unique_name(cat, gen_name(cat, gtags))
    folder = os.path.join(TPL, cat, name)
    os.makedirs(folder, exist_ok=True)

    saved = []
    used_fn = set()
    for im in images:
        fn = safe(os.path.basename(im.get("filename", "img.jpg"))) or "img.jpg"
        if not fn.lower().endswith(IMG_EXT):
            fn += ".jpg"
        stem, ext = os.path.splitext(fn)
        k = fn; n = 1
        while k in used_fn:
            k = f"{stem}-{n}{ext}"; n += 1
        used_fn.add(k)
        b64 = im.get("data_base64", "")
        if "," in b64[:64]:
            b64 = b64.split(",", 1)[1]
        with open(os.path.join(folder, k), "wb") as f:
            f.write(base64.b64decode(b64))
        rel = f"templates/{cat}/{name}/{k}"
        saved.append({"rel": rel, "tags": im.get("tags", {}),
                      "is_main": im.get("tags", {}).get("主图首图") == "是"})

    cover = next((s["rel"] for s in saved if s["is_main"]), saved[0]["rel"])
    # 唯一 group_id
    ids = existing_ids()
    base_id = "g_up_" + re.sub(r"\W+", "", name)[:24]
    gid = base_id; i = 1
    while gid in ids:
        gid = f"{base_id}_{i}"; i += 1

    append_excel(cat, gid, name, payload.get("model_name", ""),
                 payload.get("model_type", "童模") if cat == "童装" else payload.get("model_type", ""),
                 payload.get("scene_desc", ""), cover, gtags, saved)
    # 重建清单
    r = subprocess.run([sys.executable, os.path.join(HERE, "rebuild_manifest.py")],
                       capture_output=True, text=True)
    return {"ok": True, "group_name": name, "group_id": gid,
            "image_count": len(saved), "category": cat,
            "rebuild": (r.stdout or "")[-300:]}


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *a, **k):
        super().__init__(*a, directory=LIB, **k)

    def log_message(self, *a):
        pass

    def end_headers(self):
        # 禁用缓存：刷新总是拿到最新的 templates-data.js / 图片
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        super().end_headers()

    def _json(self, obj, code=200):
        body = json.dumps(obj, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):
        path = urlparse(self.path).path
        if path == "/api/add-group":
            try:
                n = int(self.headers.get("Content-Length", 0))
                payload = json.loads(self.rfile.read(n).decode("utf-8"))
                self._json(add_group(payload))
            except PermissionError:
                self._json({"ok": False, "error": "Excel 文件被占用，请先关闭 元数据录入模板.xlsx 再上传"}, 500)
            except Exception as e:
                self._json({"ok": False, "error": f"{type(e).__name__}: {e}"}, 500)
        elif path == "/api/health":
            self._json({"ok": True})
        else:
            self._json({"ok": False, "error": "未知接口"}, 404)

    def do_GET(self):
        if urlparse(self.path).path == "/api/health":
            return self._json({"ok": True, "server": True})
        return super().do_GET()


def main():
    if not os.path.exists(os.path.join(LIB, "index.html")):
        print("！没找到 index.html，请把本脚本放在 视觉模板库/工具/ 下运行。"); return
    httpd = ThreadingHTTPServer(("127.0.0.1", PORT), Handler)
    url = f"http://127.0.0.1:{PORT}/index.html"
    print("=" * 52)
    print("  视觉模板库 已启动（支持网页上传）")
    print("  打开： " + url)
    print("  右上角『+ 新增图组』即可上传图片组")
    print("  停止： 按 Ctrl+C")
    print("=" * 52)
    threading.Thread(target=lambda: (time.sleep(1), webbrowser.open(url)), daemon=True).start()
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\n已停止。")


if __name__ == "__main__":
    main()
