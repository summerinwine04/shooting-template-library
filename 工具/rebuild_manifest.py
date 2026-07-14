#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
重建清单（按类目 schema）：扫描 templates/ 图片 + 读取 元数据录入模板.xlsx，
按每个图组所属【类目】套用对应标签维度，生成 data/templates-data.js 与 .json。

用法：  python3 工具/rebuild_manifest.py
依赖：  openpyxl   (pip install openpyxl)
"""
import os, sys, json, re, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
TPL = os.path.join(ROOT, "templates")
DATA = os.path.join(ROOT, "data")
XLSX = os.path.join(ROOT, "元数据录入模板.xlsx")
DICT = os.path.join(DATA, "标签字典-分类目.json")
IMG_EXT = (".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp")

try:
    from openpyxl import load_workbook
except ImportError:
    print("缺少 openpyxl，请先运行: pip install openpyxl"); sys.exit(1)


def split_multi(v):
    if v is None:
        return []
    return [x.strip() for x in re.split(r"[、,，/\s]+", str(v)) if x.strip()]


def num(v):
    if v in (None, ""):
        return None
    try:
        f = float(v); return int(f) if f == int(f) else round(f, 4)
    except (ValueError, TypeError):
        return None


def rel(p):
    return os.path.relpath(p, ROOT).replace(os.sep, "/")


def schema_for(D, cat):
    sc = D["schemas"].get(cat)
    if sc and sc.get("group_level"):
        return cat, sc
    return "通用", D["schemas"]["通用"]


def read_excel():
    """所有 *图组 表 → {图组ID: {列名:值}}；所有 *图片 表 → {图片文件: {列名:值}}。"""
    groups, images = {}, {}
    if not os.path.exists(XLSX):
        print("  · 未找到 Excel，仅按文件夹扫描（无标签）")
        return groups, images
    wb = load_workbook(XLSX, data_only=True)
    for sn in wb.sheetnames:
        if sn.endswith("图组"):
            ws = wb[sn]; head = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                rec = dict(zip(head, row)); gid = rec.get("图组ID")
                if gid:
                    groups[str(gid).strip()] = rec
        elif sn.endswith("图片"):
            ws = wb[sn]; head = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                rec = dict(zip(head, row)); f = rec.get("图片文件")
                if f:
                    images[str(f).strip().replace("\\", "/")] = rec
    return groups, images


def metrics_of(rec):
    return {
        "production_count": num(rec.get("生产使用次数")),
        "review_pass_rate": num(rec.get("通过率")),
        "online_count": num(rec.get("上线次数")),
        "ctr": num(rec.get("CTR")),
        "sa_grade_avg": rec.get("SA档位") or None,
    }


def scan_files(D):
    found = []
    if not os.path.isdir(TPL):
        return found
    for cat in sorted(os.listdir(TPL)):
        cdir = os.path.join(TPL, cat)
        if not os.path.isdir(cdir):
            continue
        for gname in sorted(os.listdir(cdir)):
            gdir = os.path.join(cdir, gname)
            if not os.path.isdir(gdir):
                continue
            for r, _, files in os.walk(gdir):
                for fn in sorted(files):
                    if fn.lower().endswith(IMG_EXT):
                        found.append((cat, gname, rel(os.path.join(r, fn))))
    return found


def main():
    D = json.load(open(DICT, encoding="utf-8"))
    print("重建清单中（按类目 schema）…")
    g_rec, i_rec = read_excel()
    files = scan_files(D)
    print(f"  · 扫描到 {len(files)} 张图片文件")

    name2gid = {str(v.get("图组名称")): k for k, v in g_rec.items() if v.get("图组名称")}
    groups = {}
    missing = 0

    for cat, gname, relpath in files:
        gid = name2gid.get(gname) or ("auto_" + re.sub(r"\W+", "_", gname)[:40])
        sname, schema = schema_for(D, cat)
        if gid not in groups:
            rec = g_rec.get(gid, {})
            gtags = {}
            for k, m in schema["group_level"].items():
                raw = rec.get(m["label"])
                gtags[k] = split_multi(raw) if m["multi"] else (str(raw).strip() if raw not in (None, "") else None)
            groups[gid] = {
                "group_id": gid, "category": cat, "schema": sname,
                "name": rec.get("图组名称") or gname,
                "model_name": rec.get("模特") or "", "model_type": rec.get("模特类型") or "",
                "scene_desc": rec.get("场景") or "",
                "tags": gtags, "metrics": metrics_of(rec),
                "cover_file": (str(rec.get("封面文件")).strip() if rec.get("封面文件") else None),
                "images": [],
            }
        irec = i_rec.get(relpath)
        if irec is None:
            missing += 1; irec = {}
        itags = {}
        for k, m in schema["image_level"].items():
            v = irec.get(m["label"])
            if v not in (None, ""):
                itags[k] = str(v).strip()
        groups[gid]["images"].append({
            "image_id": re.sub(r"\W+", "_", relpath), "file": relpath,
            "tags": itags, "metrics": metrics_of(irec),
        })

    out = []
    for gr in groups.values():
        if not gr["images"]:
            continue
        if not gr["cover_file"]:
            gr["cover_file"] = gr["images"][0]["file"]
        gr["image_count"] = len(gr["images"])
        out.append(gr)

    order = {c: i for i, c in enumerate(D["categories"])}
    out.sort(key=lambda x: (order.get(x["category"], 99), x["name"]))

    db = {
        "meta": {
            "library_name": "AI拍摄 try-on 视觉模板库",
            "generated_at": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "group_count": len(out), "image_count": sum(len(x["images"]) for x in out),
            "categories": D["categories"],
        },
        "tag_dict": D,
        "groups": out,
    }
    os.makedirs(DATA, exist_ok=True)
    json.dump(db, open(os.path.join(DATA, "templates-data.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    with open(os.path.join(DATA, "templates-data.js"), "w", encoding="utf-8") as f:
        f.write("// 自动生成，请勿手改。运行 工具/rebuild_manifest.py 重建。\n")
        f.write("window.TEMPLATE_DB = "); json.dump(db, f, ensure_ascii=False); f.write(";\n")

    by_cat = {}
    for g in out:
        by_cat[g["category"]] = by_cat.get(g["category"], 0) + 1
    print(f"  · 图组 {db['meta']['group_count']} 个 / 图片 {db['meta']['image_count']} 张  {by_cat}")
    if missing:
        print(f"  · 其中 {missing} 张图暂无标签（按『待补全』展示）")
    print("✅ 已生成 data/templates-data.js 与 .json，刷新 index.html 即可。")


if __name__ == "__main__":
    main()
